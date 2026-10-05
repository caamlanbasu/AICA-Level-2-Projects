"""Excel writer: reads the URL list and writes results back into the same workbook.

The workbook is edited in place with openpyxl so that existing formatting is
kept. Nothing is rewritten with pandas.
"""

from __future__ import annotations

import logging
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import column_index_from_string
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.worksheet import Worksheet

from .config import (
    COMMENTS_COLUMN_WIDTH,
    COMMENTS_HEADER,
    COMMENTS_HEADER_ALIASES,
    HEADER_SCAN_ROWS,
    LOGGER_NAME,
    NEW_URL_HEADER,
    PED_COLUMN_WIDTH,
    PED_HEADER,
    URL_HEADER_KEYWORDS,
)
from .models import ExcelFormatError
from .text_utils import looks_like_url, normalise_url, url_key

LOG = logging.getLogger(LOGGER_NAME)

_MAX_SCAN_COLUMNS = 256
_HYPERLINK_FORMULA_RE = re.compile(r'^=\s*HYPERLINK\(\s*"([^"]+)"', re.IGNORECASE)


@dataclass
class UrlRow:
    """One row of the workbook that holds a usable URL."""

    row: int
    url: str


@dataclass
class RowResult:
    """What to write for one row. ``url`` guards against the sheet having changed."""

    url: str
    ped_text: Optional[str]
    comment: Optional[str]


@dataclass
class SheetLayout:
    sheet_title: str
    header_row: int          # 0 when the list has no header row at all
    url_col: int
    last_row: int


# --------------------------------------------------------------------------- #
# Cell helpers
# --------------------------------------------------------------------------- #
def _used_bounds(ws: Worksheet) -> Tuple[int, int]:
    """Last row and column that hold a value (formatting-only cells are ignored)."""
    last_row = last_col = 0
    for (row, col), cell in ws._cells.items():  # noqa: SLF001 - no public equivalent
        value = cell.value
        if value is not None and str(value).strip() != "":
            last_row = max(last_row, row)
            last_col = max(last_col, col)
    return last_row, last_col


def _text(cell: Cell) -> str:
    value = cell.value
    return "" if value is None else str(value).strip()


def _url_text(cell: Cell) -> str:
    """Cell text, or the hyperlink target when the visible text is not itself a URL."""
    text = _text(cell)
    if looks_like_url(text):
        return text
    formula = _HYPERLINK_FORMULA_RE.match(text)
    if formula:
        return formula.group(1).strip()
    link = getattr(cell, "hyperlink", None)
    target = (getattr(link, "target", None) or "").strip()
    if target.lower().startswith(("http://", "https://")):
        return target
    return text


def _header_key(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def _copy_style(source: Cell, target: Cell, include_font: bool = True) -> None:
    if not source.has_style:
        return
    if include_font:
        target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


# --------------------------------------------------------------------------- #
# Structural edits that openpyxl does not finish on its own
# --------------------------------------------------------------------------- #
def _refresh_hyperlink_refs(ws: Worksheet) -> None:
    """openpyxl moves cells on insert but leaves their hyperlink anchors behind."""
    for cell in ws._cells.values():  # noqa: SLF001
        link = getattr(cell, "hyperlink", None)
        if link is not None:
            link.ref = cell.coordinate


def _shift_ref(ref: str, idx: int, by_column: bool) -> Tuple[str, str]:
    """New range after an insert at ``idx``, and how it changed.

    Returns ``(new_ref, change)`` where change is "", "shifted" or "extended".
    """
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    change = ""
    if by_column:
        if idx <= min_col:
            min_col, max_col, change = min_col + 1, max_col + 1, "shifted"
        elif idx <= max_col:
            max_col, change = max_col + 1, "extended"
    else:
        if idx <= min_row:
            min_row, max_row, change = min_row + 1, max_row + 1, "shifted"
        elif idx <= max_row:
            max_row, change = max_row + 1, "extended"
    new_ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    return new_ref, change


def _insert_table_column(ws: Worksheet, table: object, position: int, title: str, column: int) -> None:
    """Add a column definition to an Excel table so its header list stays valid."""
    columns = list(table.tableColumns)  # type: ignore[attr-defined]
    header_row = range_boundaries(table.ref)[1]  # type: ignore[attr-defined]
    header_cell = ws.cell(row=header_row, column=column)
    existing = {str(c.name).lower() for c in columns}
    name = _text(header_cell) or title
    base, suffix = name, 2
    while name.lower() in existing:
        name = f"{base} {suffix}"
        suffix += 1
    header_cell.value = name
    new_id = max((c.id or 0 for c in columns), default=0) + 1
    columns.insert(position, TableColumn(id=new_id, name=name))
    table.tableColumns = columns  # type: ignore[attr-defined]


def _update_tables(ws: Worksheet, idx: int, header_row: int, title: str, inserted: bool) -> None:
    """Keep Excel tables and the sheet auto-filter consistent with a new column at ``idx``."""
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if inserted:
            new_ref, change = _shift_ref(table.ref, idx, by_column=True)
        else:
            new_ref, change = table.ref, ""
        adjacent = idx == max_col + 1 and min_row == header_row and change == ""
        if adjacent:
            new_ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col + 1)}{max_row}"
            change = "extended"
        if not change:
            continue
        table.ref = new_ref
        if table.autoFilter is not None and table.autoFilter.ref:
            f_min_col, f_min_row, f_max_col, f_max_row = range_boundaries(table.autoFilter.ref)
            t_min_col, _, t_max_col, _ = range_boundaries(new_ref)
            table.autoFilter.ref = (
                f"{get_column_letter(t_min_col)}{f_min_row}:{get_column_letter(t_max_col)}{f_max_row}"
            )
        if change == "extended":
            _insert_table_column(ws, table, idx - min_col, title, idx)

    if ws.auto_filter is not None and ws.auto_filter.ref:
        min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
        if inserted:
            ws.auto_filter.ref, _ = _shift_ref(ws.auto_filter.ref, idx, by_column=True)
        elif idx == max_col + 1 and min_row == header_row:
            ws.auto_filter.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(idx)}{max_row}"


def _insert_column(ws: Worksheet, idx: int) -> None:
    """Insert an empty column before ``idx``, keeping widths, merges and links in step."""
    # Merged ranges: move those to the right, widen those that straddle the new column.
    remerge: List[Tuple[int, int, int, int]] = []
    for merged in list(ws.merged_cells.ranges):
        if merged.max_col < idx:
            continue
        ws.unmerge_cells(merged.coord)
        start = merged.min_col + 1 if merged.min_col >= idx else merged.min_col
        remerge.append((merged.min_row, start, merged.max_row, merged.max_col + 1))

    # Column widths / hidden flags / column styles.
    rebuilt = []
    for letter, dimension in list(ws.column_dimensions.items()):
        start = dimension.min or column_index_from_string(letter)
        end = dimension.max or start
        if end < idx:
            rebuilt.append((start, end, dimension))
        elif start >= idx:
            rebuilt.append((start + 1, end + 1, dimension))
        else:
            rebuilt.append((start, idx - 1, dimension))
            rebuilt.append((idx + 1, end + 1, dimension))

    ws.insert_cols(idx)

    ws.column_dimensions.clear()
    for start, end, dimension in rebuilt:
        clone = copy(dimension)
        clone.index = get_column_letter(start)
        clone.min, clone.max = start, end
        ws.column_dimensions[clone.index] = clone
    for min_row, min_col, max_row, max_col in remerge:
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
    _refresh_hyperlink_refs(ws)


def _insert_top_row(ws: Worksheet) -> None:
    """Insert an empty row 1 (used when a URL list has no header row)."""
    remerge = []
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(merged.coord)
        remerge.append((merged.min_row + 1, merged.min_col, merged.max_row + 1, merged.max_col))
    heights = [(index + 1, copy(dimension)) for index, dimension in list(ws.row_dimensions.items())]

    ws.insert_rows(1)

    ws.row_dimensions.clear()
    for index, dimension in heights:
        dimension.index = index
        ws.row_dimensions[index] = dimension
    for min_row, min_col, max_row, max_col in remerge:
        ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
    _refresh_hyperlink_refs(ws)
    for table in ws.tables.values():
        table.ref, _ = _shift_ref(table.ref, 1, by_column=False)
        if table.autoFilter is not None and table.autoFilter.ref:
            table.autoFilter.ref, _ = _shift_ref(table.autoFilter.ref, 1, by_column=False)
    if ws.auto_filter is not None and ws.auto_filter.ref:
        ws.auto_filter.ref, _ = _shift_ref(ws.auto_filter.ref, 1, by_column=False)


def _set_column_width(ws: Worksheet, idx: int, width: float) -> None:
    """Set one column's width, first splitting any grouped width definition that covers it."""
    for letter, dimension in list(ws.column_dimensions.items()):
        start = dimension.min or column_index_from_string(letter)
        end = dimension.max or start
        if start == end or not start <= idx <= end:
            continue
        del ws.column_dimensions[letter]
        for piece_start, piece_end in ((start, idx - 1), (idx, idx), (idx + 1, end)):
            if piece_start > piece_end:
                continue
            clone = copy(dimension)
            clone.index = get_column_letter(piece_start)
            clone.min, clone.max = piece_start, piece_end
            ws.column_dimensions[clone.index] = clone
    ws.column_dimensions[get_column_letter(idx)].width = width


def _sheet_has_formulas(ws: Worksheet) -> bool:
    return any(cell.data_type == "f" for cell in ws._cells.values())  # noqa: SLF001


# --------------------------------------------------------------------------- #
# Public class
# --------------------------------------------------------------------------- #
class ExcelUrlList:
    """A workbook used as the URL list: read the URLs, write results back, save in place."""

    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self._workbook: Optional[Workbook] = None

    # ------------------------------------------------------------------ #
    # Layout detection
    # ------------------------------------------------------------------ #
    @staticmethod
    def _detect_layout(ws: Worksheet) -> Optional[SheetLayout]:
        """Find the URL column and header row of a sheet (``None`` if it has no URLs)."""
        last_row, last_col = _used_bounds(ws)
        if last_row == 0:
            return None
        last_col = min(last_col, _MAX_SCAN_COLUMNS)

        def url_count(col: int, first_row: int) -> int:
            return sum(1 for r in range(first_row, last_row + 1) if looks_like_url(_url_text(ws.cell(r, col))))

        reserved = {PED_HEADER.lower(), *COMMENTS_HEADER_ALIASES}

        # 1) A header that contains "url", "link" or "website".
        for row in range(1, min(last_row, HEADER_SCAN_ROWS) + 1):
            candidates = []
            for col in range(1, last_col + 1):
                cell = ws.cell(row, col)
                text = _text(cell)
                if not text or len(text) > 60 or looks_like_url(_url_text(cell)):
                    continue
                key = _header_key(text)
                if key in reserved or not any(word in key for word in URL_HEADER_KEYWORDS):
                    continue
                urls_below = url_count(col, row + 1)
                if urls_below:
                    candidates.append(("linkedin" not in key, urls_below, -col))
            if candidates:
                best = max(candidates)
                return SheetLayout(ws.title, row, -best[2], last_row)

        # 2) Otherwise the first column where most cells look like URLs.
        for col in range(1, last_col + 1):
            filled = [r for r in range(1, last_row + 1) if _text(ws.cell(r, col))]
            url_rows = [r for r in filled if looks_like_url(_url_text(ws.cell(r, col)))]
            if url_rows and len(url_rows) * 2 > len(filled):
                # The row above the first URL is the header row; row 1 URLs mean no header.
                return SheetLayout(ws.title, url_rows[0] - 1, col, last_row)
        return None

    def _locate(self, workbook: Workbook) -> Tuple[Worksheet, SheetLayout]:
        sheets = [workbook.active] + [ws for ws in workbook.worksheets if ws is not workbook.active]
        for ws in sheets:
            if ws is None or not isinstance(ws, Worksheet):
                continue
            layout = self._detect_layout(ws)
            if layout is not None:
                return ws, layout
        raise ExcelFormatError(
            "No column of website addresses was found. Add a header such as \"URL\", "
            "\"Link\" or \"Website\" above the addresses and try again."
        )

    def _open(self) -> Workbook:
        try:
            return load_workbook(self.path)
        except PermissionError:
            raise
        except Exception as exc:
            raise ExcelFormatError(f"The workbook could not be opened: {exc}") from exc

    # ------------------------------------------------------------------ #
    # Reading
    # ------------------------------------------------------------------ #
    def read_urls(self) -> List[UrlRow]:
        """Rows that hold a URL, in sheet order. Blank and invalid cells are skipped."""
        workbook = self._open()
        try:
            ws, layout = self._locate(workbook)
            rows: List[UrlRow] = []
            for row in range(layout.header_row + 1, layout.last_row + 1):
                url = normalise_url(_url_text(ws.cell(row, layout.url_col)))
                if url is not None:
                    rows.append(UrlRow(row, url))
            if not rows:
                raise ExcelFormatError("The URL column is empty.")
            LOG.debug("Excel: sheet '%s', URL column %s, header row %d, %d URL row(s)",
                      layout.sheet_title, get_column_letter(layout.url_col), layout.header_row, len(rows))
            return rows
        finally:
            workbook.close()

    # ------------------------------------------------------------------ #
    # Writing
    # ------------------------------------------------------------------ #
    @staticmethod
    def _find_header(ws: Worksheet, header_row: int, names: Tuple[str, ...]) -> Optional[int]:
        _, last_col = _used_bounds(ws)
        for col in range(1, min(last_col, _MAX_SCAN_COLUMNS) + 1):
            if _header_key(_text(ws.cell(header_row, col))) in names:
                return col
        return None

    @staticmethod
    def _column_is_empty(ws: Worksheet, col: int) -> bool:
        return not any(
            c == col and cell.value is not None and str(cell.value).strip() != ""
            for (_, c), cell in ws._cells.items()  # noqa: SLF001
        )

    def _add_column(self, ws: Worksheet, layout: SheetLayout, idx: int, title: str, width: float) -> None:
        """Create a result column at ``idx``, inserting only when that column is in use."""
        inserted = not self._column_is_empty(ws, idx)
        if inserted:
            if _sheet_has_formulas(ws):
                LOG.warning(
                    "Sheet '%s' contains formulas. Inserting the '%s' column moves cells to the right "
                    "but openpyxl does not re-point formula references; please check them.",
                    ws.title, title,
                )
            _insert_column(ws, idx)
        _update_tables(ws, idx, layout.header_row, title, inserted)

        header = ws.cell(layout.header_row, idx)
        if not _text(header):
            header.value = title
        url_header = ws.cell(layout.header_row, layout.url_col)
        if url_header.has_style:
            _copy_style(url_header, header)
        else:
            header.font = Font(bold=True)
        _set_column_width(ws, idx, width)
        for row in range(layout.header_row + 1, layout.last_row + 1):
            source = ws.cell(row, layout.url_col)
            if source.has_style:
                _copy_style(source, ws.cell(row, idx), include_font=not source.font.underline)

    def apply_results(self, results: Dict[int, RowResult]) -> None:
        """Load the workbook again and write the results into it (in memory).

        Call :meth:`save` afterwards. Rows whose URL no longer matches are skipped.
        """
        workbook = self._open()
        ws, layout = self._locate(workbook)
        offset = 0

        if layout.header_row == 0:
            _insert_top_row(ws)
            offset = 1
            layout.header_row = 1
            layout.last_row += 1
        url_header = ws.cell(layout.header_row, layout.url_col)
        if not _text(url_header):
            url_header.value = NEW_URL_HEADER
            url_header.font = Font(bold=True)

        ped_col = self._find_header(ws, layout.header_row, (PED_HEADER.lower(),))
        if ped_col is None:
            ped_col = layout.url_col + 1
            self._add_column(ws, layout, ped_col, PED_HEADER, PED_COLUMN_WIDTH)

        comments_col = self._find_header(ws, layout.header_row, COMMENTS_HEADER_ALIASES)
        if comments_col is None:
            comments_col = max(_used_bounds(ws)[1], ped_col) + 1
            self._add_column(ws, layout, comments_col, COMMENTS_HEADER, COMMENTS_COLUMN_WIDTH)

        written = 0
        for row, result in sorted(results.items()):
            target_row = row + offset
            current = normalise_url(_url_text(ws.cell(target_row, layout.url_col)))
            if current is None or url_key(current) != url_key(result.url):
                LOG.warning("Row %d no longer holds %s; result not written", target_row, result.url)
                continue
            ws.cell(target_row, ped_col).value = result.ped_text
            ws.cell(target_row, comments_col).value = result.comment
            written += 1
        LOG.debug("Excel: %d row(s) updated (Latest PED in column %s, Comments in column %s)",
                  written, get_column_letter(ped_col), get_column_letter(comments_col))
        self._workbook = workbook

    def save(self) -> None:
        """Save over the original file. Raises PermissionError while it is open in Excel."""
        if self._workbook is None:
            raise RuntimeError("apply_results() must be called before save()")
        self._workbook.save(self.path)

    def save_copy(self, target: Path) -> None:
        """Save the updated workbook under another name (used when the original stays locked)."""
        if self._workbook is None:
            raise RuntimeError("apply_results() must be called before save_copy()")
        self._workbook.save(target)
